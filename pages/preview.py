import streamlit as st
import openpyxl
import json
import re
from shared.utils import get_cell_value, resolve_placeholders
from pathlib import Path
from shared.word_generator import generate_word_from_html

__all__ = ['render_preview_tab']



def render_preview_tab(wb):
    st.title("Memo Preview")

    # Load placeholder values from config
    config_path = Path(__file__).parent.parent / "config" / "variables.json"
    with open(config_path) as f:
        placeholder_values = json.load(f)

    # Create two columns
    col1, col2 = st.columns(2)

    with col1:
        st.subheader("Input")
        input_text = st.text_area(
            "",
            height=400,
            value="""{{name}}"""
        )

    with col2:
        st.subheader("Output")
        if input_text:
            try:
                output_text = resolve_placeholders(input_text, placeholder_values, wb)
                output_html = "<div class='memo-container'>"+output_text.replace("\n", "<br>")+"</div>"
                st.markdown(
                        output_html,
                        unsafe_allow_html=True)
            except Exception as e:
                st.error(f"Error resolving placeholders: {e}")
        if st.button("📄 Generate Word Document"):
            file_path = generate_word_from_html(output_html, output_path="memo.docx")
            with open(file_path, "rb") as f:
                st.download_button(
                    "Download Word Memo",
                    f,
                    file_name="memo.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )